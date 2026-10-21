"""Build data/students.js for students.html (大学生の割合 × 国勢と住基の差).

Students: 令和2年国勢調査 就業状態等基本集計 第15-3表
  「男女，在学学校・未就学の種類別人口－全国，都道府県，市区町村」
  (e-Stat statInfId=000032201235, fileKind=0 → data/raw/students/census2020_school.xlsx)
  在学/学歴は大規模調査（2010・2020年）だけの項目で、2015年・2025年の簡易調査には無い。
  2025年も簡易調査で在学の項目が無いため、学生数は2020年の値を使う（次は2030年）。
Commuting to school: 令和2年国勢調査 従業地・通学地による人口・就業状態等集計 第1-3表
  「男女，年齢（5歳階級），常住地又は従業地・通学地別通学者数－全国，都道府県，市区町村」
  (statInfId=000032214143 → data/raw/students/commute2020.xlsx)。大学生に当たる20〜24歳を使う。
Age 20-24 share (each census, for the moving x-axis): 人口等基本集計の年齢別人口
  2010 第3-2表 各歳 (statInfId=000012460665, CSV) / 2015 第3-2表 各歳 (000031473213, CSV) /
  2020 第2-7表 5歳階級 (000032142410) / 2025 第2-7表 5歳階級・原数値 (000040506431)
  → data/raw/students/age/age{year}.(csv|xlsx)。割合＝20〜24歳÷（総数−年齢「不詳」）。
Population and 国勢−住基 differences come from data/national.json (tools/build_national.py).

Usage: python tools/build_students.py   (writes data/students.js)
"""
import csv
import json
import pathlib
import urllib.request

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "raw" / "students" / "census2020_school.xlsx"
SID = "000032201235"
# 列: 4=在学者計 8=短大・高専 9=大学 10=大学院
COL = {"junior": 8, "univ": 9, "grad": 10}


COMMUTE = ROOT / "data" / "raw" / "students" / "commute2020.xlsx"
COMMUTE_SID = "000032214143"
# 第1-3表の列: 5=常住地の通学者 12=県内他市町村へ 13=他県へ 14=通学先「不詳・外国」 15=通学地「不詳」
#              17=通学地の通学者 20=県内他市町村から 21=他県から 22=不詳等で当地に常住
AGE = "03_20～24歳"
NAME = {"13100": "特別区部"}  # 東京23区は表では「特別区部」
AGEDIR = ROOT / "data" / "raw" / "students" / "age"
AGE_SRC = {"2010": ("000012460665", 1, "csv"), "2015": ("000031473213", 1, "csv"),
           "2020": ("000032142410", 0, "xlsx"), "2025": ("000040506431", 0, "xlsx")}
CODE2010 = {"43100": "43201"}  # 熊本市 (designated city from 2012)


def download(sid, dest, kind=0):
    if dest.exists() and dest.stat().st_size > 1000:
        return
    dest.parent.mkdir(parents=True, exist_ok=True)
    url = f"https://www.e-stat.go.jp/stat-search/file-download?statInfId={sid}&fileKind={kind}"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=300) as r:
        dest.write_bytes(r.read())


def num(v):
    return int(v) if isinstance(v, (int, float)) else 0  # "-"=0


def read_commute(codes):
    """20〜24歳の通学者: [常住通学者(通学先判明), 県内他市町村へ, 他県へ, 通学地の通学者(常住地判明), 県内他市町村から, 他県から]"""
    d = {}
    ws = openpyxl.load_workbook(COMMUTE, read_only=True, data_only=True).active
    for r in ws.iter_rows(values_only=True):
        if r[3] != "0_総数" or r[4] != AGE or not r[2]:
            continue
        code = str(r[2]).split("_", 1)[0]
        if code in codes and code not in d:
            d[code] = [num(r[5]) - num(r[14]) - num(r[15]), num(r[12]), num(r[13]),
                       num(r[17]) - num(r[22]), num(r[20]), num(r[21])]
    return d


def read_age(year, codes):
    """{code: [20〜24歳, 年齢判明人口]}（国籍総数・男女総数・その年の境域の市区町村の行）"""
    sid, kind, ext = AGE_SRC[year]
    path = AGEDIR / f"age{year}.{ext}"
    download(sid, path, kind)
    d = {}
    if ext == "csv":
        head = None
        for r in csv.reader(path.read_text(encoding="cp932").splitlines()):
            if head is None:
                if "総数（年齢）" in r:
                    head = r
                    it = r.index("総数（年齢）")
                    ia = [r.index(f"{a}歳") for a in range(20, 25)]
                    iu = [i for i, h in enumerate(r) if "不詳" in h][0]
                continue
            # 大項目 01/0101＝国籍総数（2015は男女総数も込み、2010は男女が danjo.0000 の列）
            if len(r) <= iu or r[1] not in ("01", "0101") or r[4] != year or r[2] in d:
                continue
            if year == "2010" and r[6] != "danjo.0000":
                continue
            if r[2] in codes:
                n = lambda v: int(v) if v.strip().lstrip("-").isdigit() else 0
                d[r[2]] = [sum(n(r[i]) for i in ia), n(r[it]) - n(r[iu])]
    else:
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
        ia = ic = ik = None
        for r in ws.iter_rows(values_only=True):
            r = list(r)
            if ia is None and "05_20～24歳" in r:
                ia, it = r.index("05_20～24歳"), r.index("00_総数")
                iu = [i for i, h in enumerate(r) if h and "不詳" in str(h)][0]
            if ic is None and "地域識別コード" in r:  # 現在の地域コード＝最後の「…地域コード」列
                ik = r.index("地域識別コード")
                ic = [i for i, h in enumerate(r) if h and str(h).endswith("地域コード")][-1]
            if ia is None or ic is None:
                continue
            if not str(r[0]).startswith("0_国籍総数") or not str(r[1]).startswith("0_総数"):
                continue
            code = str(r[ic])
            if str(r[ik]) in ("1", "2", "3") and code in codes and code not in d:
                d[code] = [num(r[ia]), num(r[it]) - num(r[iu])]
    return d


def main():
    download(SID, SRC)
    download(COMMUTE_SID, COMMUTE)
    ws = openpyxl.load_workbook(SRC, read_only=True, data_only=True).active
    stu = {}
    for r in ws.iter_rows(values_only=True):
        if str(r[0]).startswith("0_") and str(r[1]) in ("1", "2") and r[3]:
            code, name = str(r[3]).split("_", 1)
            stu[code] = {"name": name, **{k: num(r[c]) for k, c in COL.items()}}

    nat = json.loads((ROOT / "data" / "national.json").read_text(encoding="utf-8"))
    com = read_commute({n["code"] for n in nat})
    codes = {n["code"] for n in nat} | set(CODE2010.values())
    ages = {y: read_age(y, codes) for y in AGE_SRC}
    out = []
    for n in nat:
        s = stu[n["code"]]
        if s["name"] != NAME.get(n["code"], n["city"]):
            raise SystemExit(f"市名不一致 {n['code']}: {s['name']} != {n['city']}")
        c = com[n["code"]]
        ag = [ages[y][CODE2010.get(n["code"], n["code"]) if y == "2010" else n["code"]] for y in AGE_SRC]
        out.append([n["pref"], n["city"], n["census"], n["juki"], s["univ"], s["grad"], s["junior"], c, ag])
        print(n["city"], s["univ"], s["grad"], s["junior"], c)

    js = ("// generated by tools/build_students.py — 県, 市, 国勢[2010,2015,2020,2025], "
          "住基[2011-03-31,2015/2020/2025の9月末・10月1日（未公表は翌年1/1）], 2020在学者: 大学, 大学院, 短大・高専, "
          "2020通学者20〜24歳[常住(先判明), 県内他市町村へ, 他県へ, 通学地(元判明), 県内他市町村から, 他県から], 各回国勢の[20〜24歳, 年齢判明人口](2010,2015,2020,2025)\n"
          "const STUDENTS=[\n" + ",\n".join(json.dumps(r, ensure_ascii=False, separators=(",", ":")) for r in out) + "\n];\n")
    (ROOT / "data" / "students.js").write_text(js, encoding="utf-8")
    print("wrote data/students.js")


if __name__ == "__main__":
    main()
