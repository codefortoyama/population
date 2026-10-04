"""Build towns.js / townage.js from data/raw (juki town totals + town x age).

- towns.js rows -> [label, T x7, M x7, F x7]
- townage.js -> TOWNAGE20/25 (+M/F) per bucket, None where secret
Run: python tools/build_towns.py   (after python tools/fetch_raw.py)
"""
import pathlib

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
DATA = ROOT / "data"

TOTAL_FILES = [
    ("town2009.xlsx", "2020-09"),
    ("town2109.xlsx", "2021-09"),
    ("town2209.xlsx", "2022-09"),
    ("shukei2309.xlsx", "2023-09"),
    ("town2409.xlsx", "2024-09"),
    ("town2509.xlsx", "2025-09"),
    ("town2608.xlsx", "2026-08"),
]
AGE_FILES_20 = ["ta_r2_t1.xlsx", "ta_r2_t2.xlsx", "townage_r2_6.xlsx"]
AGE_FILES_25 = ["ta_r7_t1.xlsx", "ta_r7_t2.xlsx", "ta_r7_6.xlsx"]


def parse_totals_sex(path):
    """key -> (total, male, female); None where missing/non-numeric."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    h, ord_keys, cur = {}, [], ""
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        b, c = row[1], row[2]
        if c is None or "計" in str(c):
            continue
        bn = "" if b is None else str(b).strip()
        cn = "" if c is None else str(c).strip()
        if cn in ("", "0"):
            continue
        if bn != "" and bn != "0":
            cur = bn
        town = cn
        k = cur + "|" + town
        if k not in h:
            ord_keys.append(k)
        def N(v):
            return v if isinstance(v, int) else None
        h[k] = (N(row[6]), N(row[4]), N(row[5]))
    wb.close()
    return ord_keys, h


def parse_age_sex(path):
    """key -> ([21] total or None buckets, male buckets, female buckets)."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    H, ord_keys, cur, ct = {}, [], "", ""
    def new():
        return ([None] * 21, [None] * 21, [None] * 21)
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        b, c, d = row[1], row[2], row[3]
        bn = "" if b is None else str(b).strip()
        cn = "" if c is None else str(c).strip()
        if bn != "" and bn != "0":
            cur = bn
        if cn != "" and cn != "0":
            ct = cn
        if "計" in ct or ct in ("", "0"):
            continue
        k = cur + "|" + ct
        e, f, g = row[4], row[5], row[6]
        if isinstance(g, str) and not isinstance(d, str):
            # secret total with numeric age: still record sex if numeric
            pass
        if isinstance(d, str):
            ds = d.strip()
            if "計" in ds:
                continue
            if ds.startswith("101"):
                if k not in H:
                    H[k] = new()
                    ord_keys.append(k)
                for arr, v in ((H[k][0], g), (H[k][1], e), (H[k][2], f)):
                    if isinstance(v, int):
                        arr[20] = (arr[20] or 0) + v
            continue
        if d is None:
            continue
        if k not in H:
            H[k] = new()
            ord_keys.append(k)
        bkt = min(int(d) // 5, 20)
        for arr, v in ((H[k][0], g), (H[k][1], e), (H[k][2], f)):
            if isinstance(v, int):
                arr[bkt] = (arr[bkt] or 0) + v
    wb.close()
    return ord_keys, H


def main():
    # ---- towns.js with sex ----
    per_year = {}
    orders = {}
    for fname, year in TOTAL_FILES:
        o, h = parse_totals_sex(RAW / fname)
        per_year[year] = h
        orders[year] = o
    master = list(orders["2023-09"])
    seen = set(master)
    for _, year in TOTAL_FILES:
        for k in orders[year]:
            if k not in seen:
                seen.add(k)
                master.append(k)
    years = [y for _, y in TOTAL_FILES]
    lines = ['const TOWN_YEARS=["' + '","'.join(years) + '"];', "const TOWNS=["]
    for k in master:
        lb = k.replace("|", " ").replace('"', "'")
        vals = []
        for col in (0, 1, 2):
            for y in years:
                v = per_year[y].get(k, (None, None, None))[col]
                vals.append("null" if v is None else str(v))
        lines.append('["' + lb + '",' + ",".join(vals) + "],")
    lines.append("];")
    (DATA / "towns.js").write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"towns.js: {len(master)} towns x21 cols")
    # juki city sex sums
    for y in years:
        m = sum(v[1] for v in per_year[y].values() if v[1] is not None)
        f = sum(v[2] for v in per_year[y].values() if v[1] is not None and v[2] is not None)
        print(f"  juki {y}: M={m} F={f}")

    # ---- townage with sex ----
    def merge(files):
        R, O = {}, []
        for fn in files:
            o, h = parse_age_sex(RAW / fn)
            for k in o:
                if k in R:
                    print("  OVERLAP", fn, k)
                R[k] = h[k]
                O.append(k)
        return R
    r20 = merge(AGE_FILES_20)
    r25 = merge(AGE_FILES_25)
    print(f"age keys: 20={len(r20)} 25={len(r25)}")
    outs = {"20": ([], [], []), "25": ([], [], [])}
    for k in master:
        for tag, R in (("20", r20), ("25", r25)):
            if k in R and any(v is not None for v in R[k][0]):
                for c in range(3):
                    arr = R[k][c]
                    outs[tag][c].append("null" if all(v is None for v in arr)
                                        else "[" + ",".join("null" if v is None else str(v) for v in arr) + "]")
            else:
                for c in range(3):
                    outs[tag][c].append("null")
    n20 = sum(1 for v in outs["20"][0] if v != "null")
    n25 = sum(1 for v in outs["25"][0] if v != "null")
    m20 = sum(1 for v in outs["20"][1] if v != "null")
    m25 = sum(1 for v in outs["25"][1] if v != "null")
    print(f"age total matched: 20={n20} 25={n25}; male matched: 20={m20} 25={m25}")
    out = ['const TOWNAGE_YEARS=["2020-03","2025-03"];']
    for tag, nm in (("20", "TOWNAGE20"), ("25", "TOWNAGE25")):
        out.append(f"const {nm}=[")
        out.append(",\n".join(outs[tag][0]))
        out.append("];")
    for tag, nm in (("20", "TOWNAGE20M"), ("25", "TOWNAGE25M")):
        out.append(f"const {nm}=[")
        out.append(",\n".join(outs[tag][1]))
        out.append("];")
    for tag, nm in (("20", "TOWNAGE20F"), ("25", "TOWNAGE25F")):
        out.append(f"const {nm}=[")
        out.append(",\n".join(outs[tag][2]))
        out.append("];")
    (DATA / "townage.js").write_text("\n".join(out) + "\n", encoding="utf-8")
    print("townage.js written")


if __name__ == "__main__":
    main()
